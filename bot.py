import os
import re
import openpyxl
import telebot

# ==============================
# إعدادات البوت
# ==============================

BOT_TOKEN = os.getenv("8606697797:AAF5_8IgFhl8NdIKv8TvzdgC-kn6Lj9pjec")

if not BOT_TOKEN:
    raise ValueError("لم يتم العثور على BOT_TOKEN")

bot = telebot.TeleBot(BOT_TOKEN)

EXCEL_FILE = "topeng.xlsx"


# ==============================
# تنظيف الاسم للمقارنة
# ==============================

def normalize_name(name):
    if not name:
        return ""

    name = str(name).strip()

    # إزالة المسافات الزائدة
    name = re.sub(r"\s+", " ", name)

    # توحيد بعض أشكال الأحرف العربية
    name = name.replace("أ", "ا")
    name = name.replace("إ", "ا")
    name = name.replace("آ", "ا")
    name = name.replace("ى", "ي")

    return name.casefold()


# ==============================
# قراءة ملف Excel
# ==============================

def find_student(student_name):

    if not os.path.exists(EXCEL_FILE):
        return {
            "type": "error",
            "message": "حدث خطأ: ملف البيانات غير موجود."
        }

    try:
        workbook = openpyxl.load_workbook(
            EXCEL_FILE,
            read_only=True,
            data_only=True
        )

        # ==========================
        # ورقة S1
        # ==========================

        if "S1" not in workbook.sheetnames:
            return {
                "type": "error",
                "message": "حدث خطأ: لم يتم العثور على صفحة S1."
            }

        sheet_s1 = workbook["S1"]

        searched_name = normalize_name(student_name)

        # --------------------------------
        # أولاً: البحث في العمود B
        # --------------------------------

        for row in range(1, sheet_s1.max_row + 1):

            name_b = sheet_s1.cell(row=row, column=2).value

            if normalize_name(name_b) == searched_name:

                alternative_name = sheet_s1.cell(row=row, column=5).value
                alternative_number = sheet_s1.cell(row=row, column=6).value

                workbook.close()

                return {
                    "type": "alternative",
                    "name": alternative_name,
                    "number": alternative_number
                }

        # --------------------------------
        # ثانياً: البحث في العمود E
        # --------------------------------

        for row in range(1, sheet_s1.max_row + 1):

            name_e = sheet_s1.cell(row=row, column=5).value

            if normalize_name(name_e) == searched_name:

                student_name_b = sheet_s1.cell(row=row, column=2).value
                student_number_c = sheet_s1.cell(row=row, column=3).value

                workbook.close()

                return {
                    "type": "alternative_reverse",
                    "name": student_name_b,
                    "number": student_number_c
                }

        # ==========================
        # ورقة Form1
        # ==========================

        if "Form1" not in workbook.sheetnames:
            workbook.close()

            return {
                "type": "not_found"
            }

        sheet_form1 = workbook["Form1"]

        # البحث في العمود D
        for row in range(1, sheet_form1.max_row + 1):

            name_d = sheet_form1.cell(row=row, column=4).value

            if normalize_name(name_d) == searched_name:

                workbook.close()

                return {
                    "type": "no_alternative"
                }

        workbook.close()

        # لم يتم العثور على الاسم بأي مكان
        return {
            "type": "name_not_found"
        }

    except Exception as e:

        return {
            "type": "error",
            "message": f"حدث خطأ أثناء قراءة ملف البيانات:\n{e}"
        }


# ==============================
# /start
# ==============================

@bot.message_handler(commands=["start"])
def start(message):

    bot.send_message(
        message.chat.id,
        "أهلاً بك 👋\n\n"
        "أرسل اسمك لإرسال البديل الخاص بك."
    )


# ==============================
# استقبال اسم الطالب
# ==============================

@bot.message_handler(func=lambda message: True)
def receive_name(message):

    student_name = message.text.strip()

    if not student_name:

        bot.send_message(
            message.chat.id,
            "يرجى إرسال اسم الطالب."
        )

        return

    # رسالة انتظار
    waiting_message = bot.send_message(
        message.chat.id,
        "🔎 جاري البحث، يرجى الانتظار..."
    )

    result = find_student(student_name)

    # حذف رسالة الانتظار
    try:
        bot.delete_message(
            message.chat.id,
            waiting_message.message_id
        )
    except:
        pass

    # ==============================
    # يوجد بديل
    # ==============================

    if result["type"] == "alternative":

        name = result["name"]
        number = result["number"]

        bot.send_message(
            message.chat.id,
            "✅ تم العثور على البديل\n\n"
            f"👤 اسم البديل: {name}\n"
            f"📱 الرقم: {number}"
        )

        return

    # ==============================
    # الطالب موجود كبديل لشخص آخر
    # ==============================

    if result["type"] == "alternative_reverse":

        name = result["name"]
        number = result["number"]

        bot.send_message(
            message.chat.id,
            "✅ تم العثور على الطالب المرتبط بك\n\n"
            f"👤 اسم الطالب: {name}\n"
            f"📱 الرقم: {number}"
        )

        return

    # ==============================
    # موجود في Form1 ولكن لا يوجد بديل
    # ==============================

    if result["type"] == "no_alternative":

        bot.send_message(
            message.chat.id,
            "⚠️ لا يتم العثور على بديل بعد."
        )

        return

    # ==============================
    # الاسم غير موجود
    # ==============================

    if result["type"] == "name_not_found":

        bot.send_message(
            message.chat.id,
            "❌ تأكد من كتابة الاسم بالشكل الصحيح.\n\n"
            "لم يتم العثور على الاسم."
        )

        return

    # ==============================
    # خطأ
    # ==============================

    if result["type"] == "error":

        bot.send_message(
            message.chat.id,
            f"⚠️ {result['message']}"
        )

        return


# ==============================
# تشغيل البوت
# ==============================

print("TOP ENG Bot is running...")

bot.infinity_polling(
    skip_pending=True,
    timeout=60,
    long_polling_timeout=60
)